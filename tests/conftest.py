import pytest
from pathlib import Path

@pytest.fixture(scope="module")
def fixtures_dir(tmp_path_factory):
    d = tmp_path_factory.mktemp("fixtures")
    
    # 1. Create sample.md
    md_file = d / "sample.md"
    md_file.write_text("# Introduction\nThis is a test markdown file. It has some text.\n\n## Section 1\nMore text here.")
    
    # 2. Create sample.txt
    txt_file = d / "sample.txt"
    txt_file.write_text("This is plain text. Nothing fancy here.")
    
    # 3. Create sample.pdf using fitz
    import fitz
    pdf_doc = fitz.open()
    page1 = pdf_doc.new_page()
    page1.insert_text((50, 50), "First Heading\nThis is page 1. It contains some text.")
    page2 = pdf_doc.new_page()
    page2.insert_text((50, 50), "Second Heading\nThis is page 2. It has more text. This should be long enough to avoid OCR.")
    pdf_file = d / "sample.pdf"
    pdf_doc.save(str(pdf_file))
    pdf_doc.close()
    
    # 4. Create sample.docx using python-docx
    import docx
    docx_doc = docx.Document()
    docx_doc.add_heading("Docx Heading", level=1)
    docx_doc.add_paragraph("This is a paragraph in docx.")
    docx_file = d / "sample.docx"
    docx_doc.save(str(docx_file))
    
    return d
